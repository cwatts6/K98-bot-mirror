"""Typed physical exports and ID fencing; no live provider operations."""

from copy import deepcopy
from dataclasses import replace
from io import BytesIO
import json
import re
from uuid import uuid4
import zipfile

import pytest

from kvk.dal.new_source_import_dal import SourceConflict
from kvk.rendering.new_source_workbook import (
    DIRECTORY,
    alias,
    build_workbook_generation,
    excel_bundle,
    partition_workbook,
)
from kvk.services.new_source_delivery_service import deliver_export
from kvk.services.new_source_workbook_transport import ReaderSheetsTransport
from services.export_execution_protocol import ProtocolError, from_http
from tests.test_kvk_source_checkpoints import build, fixture
from tests.test_kvk_source_delivery import GoogleMemoryAPI, google_delivery


class ReaderAPI(GoogleMemoryAPI):
    def __init__(self):
        super().__init__()
        self.cells, self.filters = {}, {}
        self.fail_format = False
        self.delayed = []
        self.hyperlink_readback = False
        self.widths = {}

    def execute(self, path, method, args, retries):
        file_id = args.get("spreadsheetId")
        if path == "/spreadsheets" and method == "batchUpdate":
            requests = args["body"]["requests"]
            if requests and all("updateDimensionProperties" in r for r in requests):
                self.calls.append((path, method, deepcopy(args)))
                for r in requests:
                    d = r["updateDimensionProperties"]
                    grid = d["range"]
                    if grid["sheetId"] not in {p["sheetId"] for p in self.grids[file_id].values()}:
                        raise ValueError("Retired immutable grid")
                    for i in range(grid["startIndex"], grid["endIndex"]):
                        self.widths[file_id, grid["sheetId"], i] = d["properties"]["pixelSize"]
                return {}
            if len(requests) == 1 and (
                "updateCells" in requests[0] or "setBasicFilter" in requests[0]
            ):
                self.calls.append((path, method, deepcopy(args)))
                if "setBasicFilter" in requests[0]:
                    grid = requests[0]["setBasicFilter"]["filter"]["range"]
                    if self.fail_format:
                        raise TimeoutError("format response lost")
                    self.filters[file_id, grid["sheetId"]] = deepcopy(grid)
                    return {}
                instruction = requests[0]["updateCells"]
                start = instruction["start"]
                ids = {p["sheetId"]: name for name, p in self.grids[file_id].items()}
                if start["sheetId"] not in ids:
                    raise ValueError("Retired immutable grid")
                name = ids[start["sheetId"]]
                rows = self.cells.setdefault((file_id, start["sheetId"]), [])
                offset = start["rowIndex"]
                incoming = [deepcopy(r["values"]) for r in instruction["rows"]]
                while len(rows) < offset + len(incoming):
                    rows.append([])
                rows[offset : offset + len(incoming)] = incoming
                self.values[file_id, name] = [
                    [
                        c.get("userEnteredValue", {}).get(
                            "stringValue", c.get("userEnteredValue", {}).get("numberValue", "")
                        )
                        for c in row
                    ]
                    for row in rows
                ]
                if self.fail_pointer and name == "Sheet1":
                    raise TimeoutError("pointer response lost")
                if self.fail_write and name != "Sheet1":
                    raise TimeoutError("data response lost")
                return {}
        if path == "/spreadsheets" and method == "get" and args.get("ranges"):
            self.calls.append((path, method, deepcopy(args)))
            match = re.fullmatch(r"'((?:[^']|'')+)'!A(\d+):[A-Z]+(\d+)", args["ranges"][0])
            name, start, end = match.groups()
            name, start, end = name.replace("''", "'"), int(start) - 1, int(end)
            prop = self.grids[file_id][name]
            rows = deepcopy(self.cells.get((file_id, prop["sheetId"]), [])[start:end])
            # Google omits blank ExtendedValue objects and trailing empty cells.
            for row in rows:
                for cell in row:
                    if self.hyperlink_readback:
                        fmt = cell.get("userEnteredFormat", {})
                        if "backgroundColor" in fmt:
                            fmt["backgroundColor"] = {
                                k: round(int(v * 255) / 255, 8)
                                for k, v in fmt["backgroundColor"].items()
                            }
                            fmt["backgroundColorStyle"] = {
                                "rgbColor": deepcopy(fmt["backgroundColor"])
                            }
                        text = fmt.get("textFormat", {})
                        if "link" in text:
                            cell["hyperlink"] = text.pop("link")["uri"]
                            if not text:
                                fmt.pop("textFormat")
                            if not fmt:
                                cell.pop("userEnteredFormat", None)
                    if cell.get("userEnteredValue") == {}:
                        cell.pop("userEnteredValue")
                while row and not row[-1]:
                    row.pop()
            return dict(
                sheets=[
                    dict(
                        properties=deepcopy(prop),
                        data=[
                            dict(
                                startRow=start,
                                rowData=[dict(values=r) for r in rows],
                                columnMetadata=[
                                    dict(
                                        pixelSize=self.widths.get(
                                            (file_id, prop["sheetId"], i), 100
                                        )
                                    )
                                    for i in range(prop["gridProperties"]["columnCount"])
                                ],
                            )
                        ],
                        basicFilter=dict(range=self.filters.get((file_id, prop["sheetId"]), {})),
                    )
                ]
            )
        return super().execute(path, method, args, retries)


def setup_reader(*, names=("Pass 4",), skip=()):
    args, repo, old = google_delivery()
    api = ReaderAPI()
    api.files_data, api.grids = deepcopy(old.files_data), deepcopy(old.grids)
    previous = args["transport"]
    transport = ReaderSheetsTransport(
        drive=api.client(),
        sheets=api.client(),
        registration=previous.registration,
        reuse_guard=lambda *_: False,
        protected_file_ids=(),
    )
    generation = build_workbook_generation(build(fixture(names=names, skip=skip)))
    args.update(generation=generation, selection=generation.selections[0], transport=transport)
    repo.selected = generation.selections[0]
    return args, repo, api


def test_reader_delivery_typed_values_headers_ids_and_directory():
    args, _, api = setup_reader()
    result = deliver_export(**args)
    assert result.delivery_state == "confirmed", result
    generation, transport = args["generation"], args["transport"]
    manifest = generation.manifest()
    assert transport.verify(args["destination"], generation.key) == manifest
    grids = api.grids["fake-1"]
    assert next(iter(grids)) == DIRECTORY
    assert len(grids) == 39
    player = next(t for t in generation.tables if t.name == "Pass 4 Players")
    rows = api.cells["fake-1", grids[player.name]["sheetId"]]
    assert rows[5][0]["userEnteredValue"] == {"stringValue": "Governor ID"}
    assert rows[6][0]["userEnteredValue"] == {"stringValue": "1001"}
    assert rows[6][7]["userEnteredValue"] == {"numberValue": 50_000_000}
    assert rows[6][14]["userEnteredFormat"]["numberFormat"]["pattern"] == "0.00%"
    assert grids[player.name]["gridProperties"]["frozenColumnCount"] == 2
    assert api.widths["fake-1", grids[player.name]["sheetId"], 1] == 240
    assert api.widths["fake-1", grids[DIRECTORY]["sheetId"], 0] == 340
    assert rows[6][1]["userEnteredFormat"]["wrapStrategy"] == "WRAP"
    assert "wrapStrategy" not in rows[6][7]["userEnteredFormat"]
    assert all(len(title) <= 31 for title in grids)
    assert api.values["fake-index", "Sheet1"][0][:2] == [generation.key, "1"]
    assert not any(
        path == "/spreadsheets/values" and method == "update" for path, method, _ in api.calls
    )


def test_google_whole_cell_hyperlinks_verify_on_cold_readback():
    args, _, api = setup_reader()
    api.hyperlink_readback = True
    assert deliver_export(**args).delivery_state == "confirmed"
    old = args["transport"]
    cold = ReaderSheetsTransport(
        drive=api.client(),
        sheets=api.client(),
        registration=old.registration,
        reuse_guard=lambda *_: False,
        protected_file_ids=(),
    )
    assert cold.verify(args["destination"], args["generation"].key) == args["generation"].manifest()


@pytest.mark.parametrize(
    "damage", ["foreign", "conflicting", "partial", "formula", "color", "style"]
)
def test_google_hyperlink_readback_rejects_changed_targets_and_runs(damage):
    args, _, api = setup_reader()
    assert deliver_export(**args).delivery_state == "confirmed"
    cell = api.cells["fake-1", api.grids["fake-1"][DIRECTORY]["sheetId"]][6][4]
    if damage == "foreign":
        api.hyperlink_readback = True
        cell["userEnteredFormat"]["textFormat"]["link"]["uri"] = "https://example.invalid"
    elif damage == "conflicting":
        cell["hyperlink"] = "https://example.invalid"
    elif damage == "partial":
        cell["textFormatRuns"] = [
            {"startIndex": 1, "format": {"link": {"uri": "https://example.invalid"}}}
        ]
    elif damage == "formula":
        api.hyperlink_readback = True
        cell["userEnteredValue"] = {"formulaValue": '=HYPERLINK("https://example.invalid")'}
    else:
        header = api.cells["fake-1", api.grids["fake-1"]["Report Guide"]["sheetId"]][5][0]
        if damage == "color":
            api.hyperlink_readback = True
            header["userEnteredFormat"]["backgroundColor"]["red"] = 0.2
        else:
            header["userEnteredFormat"]["backgroundColorStyle"] = {"themeColor": "ACCENT1"}
    with pytest.raises(SourceConflict):
        args["transport"].verify(args["destination"], args["generation"].key)


def test_excel_bundle_roundtrip_and_no_formulas():
    from openpyxl import load_workbook

    generation = build_workbook_generation(build(fixture()))
    content = excel_bundle(generation)
    with zipfile.ZipFile(BytesIO(content)) as archive:
        book = load_workbook(BytesIO(archive.read("report-01.xlsx")))
        assert book.sheetnames[0] == DIRECTORY
        assert len(book.sheetnames) == 63
        sheet = book["Pass 4 Players"]
        assert sheet["A6"].value == "Governor ID"
        assert sheet["A7"].data_type == "s"
        assert sheet["H7"].value == 50_000_000 and sheet["H7"].data_type == "n"
        assert sheet["H7"].number_format == "#,##0"
        assert sheet["E7"].number_format == "#,##0"
        assert sheet.column_dimensions["B"].width > sheet.column_dimensions["H"].width
        assert sheet["B7"].alignment.wrap_text is True
        assert sheet["H7"].alignment.wrap_text is not True
        assert book[DIRECTORY].column_dimensions["A"].width > 45
        assert book[DIRECTORY]["A14"].alignment.wrap_text is True
        assert sheet["O7"].number_format == "0.00%"
        assert sheet.freeze_panes == "C7" and sheet.auto_filter.ref.startswith("A6:")
        assert all(c.data_type != "f" for ws in book for row in ws for c in row)
        assert json.loads(archive.read("manifest.json"))["key"] == generation.key
        assert "reader-exact.json" in archive.namelist() and "audit.json" in archive.namelist()


def test_optional_decimals_keep_fractions_without_integer_decimal_points():
    from decimal import Decimal

    projection = build(fixture(names=("Pass 4",)))
    player = projection.sheets[1]
    rows, exact = [list(r) for r in player.rows], [list(r) for r in player.exact_rows]
    rows[0][5], exact[0][5] = Decimal("1234.125"), "1234.125"
    player = replace(player, rows=tuple(map(tuple, rows)), exact_rows=tuple(map(tuple, exact)))
    generation = build_workbook_generation(
        replace(projection, sheets=(projection.sheets[0], player, *projection.sheets[2:]))
    )
    cells = next(t for t in generation.tables if t.name == player.title).raw_values()[6]
    assert cells[4]["userEnteredFormat"]["numberFormat"]["pattern"] == "#,##0"
    assert cells[5]["userEnteredFormat"]["numberFormat"]["pattern"] == "#,##0.######"
    assert cells[5]["userEnteredValue"] == {"numberValue": 1234.125}
    assert cells[6]["userEnteredFormat"]["numberFormat"]["pattern"] == '0.00"×"'
    assert cells[14]["userEnteredFormat"]["numberFormat"]["pattern"] == "0.00%"


@pytest.mark.parametrize(
    "text", ["=1+1", "+SUM(A1)", "-1+2", "@SUM(A1)", ' \t=HYPERLINK("bad")', "玩家 😊 ملف.xlsx"]
)
def test_untrusted_text_roundtrips_literally(text):
    from openpyxl import load_workbook

    projection = build(fixture(names=("Pass 4",)))
    player = projection.sheets[1]
    rows, exact = [list(r) for r in player.rows], [list(r) for r in player.exact_rows]
    rows[0][1] = exact[0][1] = text
    player = replace(player, rows=tuple(map(tuple, rows)), exact_rows=tuple(map(tuple, exact)))
    generation = build_workbook_generation(
        replace(projection, sheets=(projection.sheets[0], player, *projection.sheets[2:]))
    )
    with zipfile.ZipFile(BytesIO(excel_bundle(generation))) as archive:
        book = load_workbook(BytesIO(archive.read("report-01.xlsx")))
        assert book[player.title]["B7"].value == text
        assert book[player.title]["B7"].data_type == "s"
    args, _, api = setup_reader()
    args["generation"] = generation
    assert deliver_export(**args).delivery_state == "confirmed"
    sid = api.grids["fake-1"][player.title]["sheetId"]
    assert api.cells["fake-1", sid][6][1]["userEnteredValue"] == {"stringValue": text}


def test_aliases_collision_unicode_and_windows_invalid_characters():
    used = set()
    labels = [
        "very long report " * 4,
        "very long report " * 4,
        "VERY LONG REPORT " * 4,
        "History",
        "'[]:/?*\\'",
        "玩家 😊",
    ]
    names = [alias(t, used) for t in labels]
    assert len({n.casefold() for n in names}) == len(names)
    assert all(0 < len(n) <= 31 and not re.search(r"[\[\]:*?/\\]", n) for n in names)


@pytest.mark.parametrize(
    "names,skip,count",
    [
        (("Pass 4",), (), 38),
        (("Pass 4", "Altar 1", "Altar 2", "Altar 3", "Pass 7"), (), 62),
        (("Pass 4", "Altar 1", "Altar 2", "Altar 3", "Pass 7"), (2, 3), 50),
    ],
)
def test_partition_budget_and_complete_pairs(names, skip, count):
    generation = build_workbook_generation(build(fixture(names=names, skip=skip)))
    manifest = generation.manifest()
    assert len(manifest) == count
    parts = partition_workbook(manifest, max_cells=11_000)
    assert parts == partition_workbook(manifest, max_cells=11_000)
    assert len(parts) > 1
    for i, pieces in enumerate(parts):
        cells = sum(max(7, p["rows"] + 6) * p["columns"] for p in pieces)
        if i == 0:
            cells += (sum(map(len, parts)) + 6) * 5
        assert cells <= 11_000
    for checkpoint in generation.projection.checkpoints:
        assert set(checkpoint.sheet_keys) <= {
            manifest[p["table"]]["logical_key"] for part in parts for p in part
        }
    for name, spec in manifest.items():
        assert sum(p["rows"] for part in parts for p in part if p["table"] == name) == spec["rows"]


@pytest.mark.parametrize("failure", ["fail_write", "fail_format", "fail_pointer"])
def test_uncertain_mutations_hold_claim_and_do_not_replay(failure):
    args, repo, api = setup_reader()
    setattr(api, failure, True)
    result = deliver_export(**args)
    assert result.delivery_state == "uncertain"
    calls = len(api.calls)
    with pytest.raises(SourceConflict):
        deliver_export(**args)
    assert len(api.calls) == calls
    if failure != "fail_pointer":
        assert not api.values.get(("fake-index", "Sheet1"))


def test_cold_readback_and_wrong_generation_payload():
    args, _, api = setup_reader()
    assert deliver_export(**args).delivery_state == "confirmed"
    old = args["transport"]
    fresh = ReaderSheetsTransport(
        drive=api.client(),
        sheets=api.client(),
        registration=old.registration,
        reuse_guard=lambda *_: False,
        protected_file_ids=(),
    )
    generation = args["generation"]
    assert fresh.verify(args["destination"], generation.key) == generation.manifest()
    values = generation.tables[1].raw_values()
    values[6][1]["userEnteredValue"] = {"stringValue": "Different"}
    with pytest.raises(SourceConflict, match="immutable"):
        fresh.write_range(
            args["destination"],
            generation.key,
            generation.tables[1].name,
            "A1",
            values,
            value_input_option="RAW",
        )


def test_retirement_fences_late_data_and_index_writes_by_id():
    args, _, api = setup_reader()
    assert deliver_export(**args).delivery_state == "confirmed"
    transport = args["transport"]
    writes = [
        deepcopy(a)
        for path, method, a in api.calls
        if method == "batchUpdate" and "updateCells" in a["body"]["requests"][0]
    ]
    transport._request_guard = lambda *_: None
    for file_id in ("fake-1", "fake-index"):
        transport.rollover_clear(file_id)
    before = deepcopy(api.values)
    for request in (writes[0], writes[-1]):
        with pytest.raises(ValueError, match="Retired"):
            api.execute("/spreadsheets", "batchUpdate", request, 0)
    assert api.values == before


def test_recorded_sdk_accepts_only_inert_typed_requests():
    from services.export_provider_adapter import recorded_clients

    sheets, _ = recorded_clients()
    args, _, api = setup_reader()
    assert deliver_export(**args).delivery_state == "confirmed"
    for path, method, arguments in api.calls:
        if path != "/spreadsheets" or method not in {"get", "batchUpdate"}:
            continue
        request = getattr(sheets.spreadsheets(), method)(**arguments)
        from_http(
            method=request.method,
            uri=request.uri,
            body=request.body,
            stream_id=str(uuid4()),
            request_id=str(uuid4()),
        )


@pytest.mark.parametrize(
    "damage",
    ["number_as_text", "formula", "filter", "grid_id", "extra_tab", "pane", "link", "width"],
)
def test_readback_rejects_type_structure_and_navigation_damage(damage):
    args, _, api = setup_reader()
    assert deliver_export(**args).delivery_state == "confirmed"
    prop = api.grids["fake-1"]["Pass 4 Players"]
    rows = api.cells["fake-1", prop["sheetId"]]
    if damage == "number_as_text":
        rows[6][7]["userEnteredValue"] = {"stringValue": "50000000"}
    elif damage == "formula":
        rows[6][7]["userEnteredValue"] = {"formulaValue": "=50000000"}
    elif damage == "filter":
        api.filters["fake-1", prop["sheetId"]]["endRowIndex"] -= 1
    elif damage == "grid_id":
        prop["sheetId"] += 1000
    elif damage == "extra_tab":
        api.grids["fake-1"]["Stale"] = dict(deepcopy(prop), title="Stale", sheetId=9999)
    elif damage == "pane":
        prop["gridProperties"]["frozenRowCount"] = 1
    elif damage == "width":
        api.widths["fake-1", prop["sheetId"], 1] = 100
    else:
        rows[6][1].setdefault("userEnteredFormat", {})["textFormat"] = {
            "link": {"uri": "https://example.invalid"}
        }
    with pytest.raises(SourceConflict):
        args["transport"].verify(args["destination"], args["generation"].key)


@pytest.mark.parametrize("public,repair", [(False, False), (True, False), (False, True)])
def test_coordinator_persists_ids_before_writes_and_pins_retry_history(monkeypatch, public, repair):
    from threading import Event
    from types import SimpleNamespace
    from unittest.mock import Mock

    from kvk.dal.new_source_import_dal import digest
    from kvk.services import new_source_reader_service
    from kvk.services.new_source_delivery_service import deliver_coordinated_export

    args, repo, api = setup_reader()
    transport, original = args["transport"], args["generation"]
    if public:
        transport.registration = replace(transport.registration, audience="public_viewer")
    loader = Mock(return_value=original.projection)
    monkeypatch.setattr(new_source_reader_service, "load_reader_projection", loader)
    job = dict(
        ConsumerKind="new_source",
        IntentID="intent-a",
        InputHash=b"a" * 32,
        KVK_NO=16,
        DestinationSetHash=digest(tuple(sorted(("fake-index", "fake-1", "fake-2")))),
    )
    if repair:
        job["RepairID"] = str(uuid4())
    claim, dal, retained = SimpleNamespace(fence=10), Mock(), {}

    def begin(_, manifest, parts):
        assert not any(
            method in {"update", "delete", "batchUpdate", "create"} for _, method, _ in api.calls
        )
        assert manifest["reader_index_id"] == 0
        assert len(manifest["reader_bindings"]["fake-1"]) == 39
        assert len(json.dumps(dict(generation=manifest, parts=parts)).encode("utf-16-le")) <= 65536
        retained.update(manifest=deepcopy(manifest), parts=deepcopy(parts))
        # New imports/selections cannot cause the captured projection to be rebuilt.
        repo.selected = replace(repo.selected, selection_version=99)
        return "attempt-a"

    dal.begin_attempt.side_effect = begin
    receipt = deliver_coordinated_export(
        job=job,
        claim=claim,
        dal=dal,
        transport=transport,
        connect=None,
        budget=lambda: None,
        stop=Event(),
        reader_guild_id="123",
    )
    loader.assert_called_once()
    assert retained["manifest"]["content_key"] == original.key
    assert retained["manifest"]["tables"] == original.manifest()
    assert receipt["files"] == ["fake-index", "fake-1"]
    dal.confirm.assert_called_once()


def test_matching_published_generation_is_private_before_retry_writes():
    args, _, api = setup_reader()
    transport, generation = args["transport"], args["generation"]
    transport.registration = replace(transport.registration, audience="public_viewer")
    assert deliver_export(**args).delivery_state == "confirmed"
    assert any(p["type"] == "anyone" for p in api.files_data["fake-1"]["permissions"])
    start = len(api.calls)
    transport.ensure_private(args["destination"], generation.key, generation.manifest())
    assert not any(p["type"] == "anyone" for p in api.files_data["fake-1"]["permissions"])
    calls = api.calls[start:]
    revoke = next(
        i
        for i, (path, method, _) in enumerate(calls)
        if path == "/permissions" and method == "delete"
    )
    write = next(
        i
        for i, (path, method, _) in enumerate(calls)
        if path == "/spreadsheets" and method == "batchUpdate"
    )
    assert revoke < write


def test_retained_export_survives_second_generation_and_refuses_slot_reuse():
    from kvk.services.new_source_delivery_service import DestinationSetupRequired

    args, _, api = setup_reader()
    transport = args["transport"]
    assert deliver_export(**args).delivery_state == "confirmed"
    original = deepcopy(api.cells)
    second = build_workbook_generation(build(fixture(names=("Pass 4", "1st Altar"))))
    transport.prepare_generation(second)
    transport.ensure_private(args["destination"], second.key, second.manifest())
    for table in second.tables:
        transport.write_range(
            args["destination"],
            second.key,
            table.name,
            "A1",
            table.raw_values(),
            value_input_option="RAW",
        )
    assert transport.verify(args["destination"], second.key) == second.manifest()
    assert all(api.cells[k] == v for k, v in original.items())
    third = build_workbook_generation(build(fixture()))
    transport.prepare_generation(third)
    with pytest.raises(DestinationSetupRequired):
        transport.ensure_private(args["destination"], third.key, third.manifest())


def test_excel_report_title_is_literal_and_oversize_text_is_rejected():
    from openpyxl import load_workbook

    from kvk.rendering.new_source_workbook import cell_value

    generation = build_workbook_generation(build(fixture(names=("=1+1",))))
    with zipfile.ZipFile(BytesIO(excel_bundle(generation))) as archive:
        book = load_workbook(BytesIO(archive.read("report-01.xlsx")))
        assert all(c.data_type != "f" for ws in book for row in ws for c in row)
    with pytest.raises(ValueError, match="lossless"):
        cell_value("x" * 32768)


def test_row_partition_headers_are_repeated_with_exact_row_coverage():
    generation = build_workbook_generation(build(fixture(names=("Pass 4",))))
    parts = partition_workbook(generation.manifest(), max_rows=8)
    pieces = [p for part in parts for p in part if p["table"] == "Pass 4 Players"]
    assert len(pieces) == 2
    assert [(p["offset"], p["rows"]) for p in pieces] == [(0, 2), (2, 2)]


def test_authority_rejects_formula_extended_value():
    from services.export_provider_adapter import recorded_clients

    sheets, _ = recorded_clients()
    body = {
        "requests": [
            {
                "updateCells": {
                    "start": {"sheetId": 1, "rowIndex": 0, "columnIndex": 0},
                    "rows": [{"values": [{"userEnteredValue": {"formulaValue": "=1+1"}}]}],
                    "fields": "userEnteredValue",
                }
            }
        ]
    }
    request = sheets.spreadsheets().batchUpdate(spreadsheetId="fake-1", body=body)
    with pytest.raises(ProtocolError):
        from_http(
            method=request.method,
            uri=request.uri,
            body=request.body,
            stream_id=str(uuid4()),
            request_id=str(uuid4()),
        )


def test_reader_publication_timeout_can_reconcile_after_restart_without_resend():
    from types import SimpleNamespace

    args, _, api = setup_reader()
    api.fail_pointer = True
    result = deliver_export(**args)
    assert result.delivery_state == "uncertain" and result.export_complete
    old = args["transport"]
    fresh = ReaderSheetsTransport(
        drive=api.client(),
        sheets=api.client(),
        registration=old.registration,
        reuse_guard=lambda *_: False,
        protected_file_ids=(),
    )
    writes_before = sum(method == "batchUpdate" for _, method, _ in api.calls)
    probe = SimpleNamespace(receipt=result.receipt, fence=1, selection=args["selection"])
    state, receipt = fresh.reconcile(
        args["destination"], probe, pinned_manifest=args["generation"].manifest()
    )
    assert state == "confirmed" and json.loads(receipt)["export_complete"]
    assert sum(method == "batchUpdate" for _, method, _ in api.calls) == writes_before


def test_stale_owner_is_rejected_before_first_reader_mutation():
    args, _, api = setup_reader()

    def guard(request):
        if request.method != "GET":
            raise SourceConflict("Stale owner")

    args["transport"]._request_guard = guard
    result = deliver_export(**args)
    assert result.delivery_state == "uncertain"
    assert not any(
        method in {"update", "batchUpdate", "delete", "create"} for _, method, _ in api.calls
    )


@pytest.mark.parametrize(
    "bad",
    [
        {"properties": {"pixelSize": True}},
        {"properties": {"pixelSize": 10000}},
        {"properties": {"pixelSize": 240, "hiddenByUser": True}},
        {"fields": "*"},
        {"range": {"sheetId": 1, "dimension": "ROWS", "startIndex": 0, "endIndex": 1}},
        {"range": {"sheetId": 1, "dimension": "COLUMNS", "startIndex": -1, "endIndex": 1}},
        {"range": {"sheetId": 1, "dimension": "COLUMNS", "startIndex": 0, "endIndex": 16385}},
    ],
)
def test_column_width_authority_rejects_unbounded_or_unrelated_mutations(bad):
    from services.export_reader_protocol import validate_reader_batch

    instruction = dict(
        range=dict(sheetId=1, dimension="COLUMNS", startIndex=0, endIndex=1),
        properties=dict(pixelSize=240),
        fields="pixelSize",
    )
    validate_reader_batch("updateDimensionProperties", instruction)
    instruction.update(bad)
    with pytest.raises(ProtocolError):
        validate_reader_batch("updateDimensionProperties", instruction)
