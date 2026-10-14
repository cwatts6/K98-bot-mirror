"""Versioned physical reader workbooks, shared by Sheets and local Excel output.

Extended values are explicit: imported text can never become a formula. The
logical projection and its exact audit remain the authority for decimal values.
"""

from dataclasses import dataclass
from datetime import UTC, datetime
from decimal import Decimal
import hashlib
import io
import json
import re
import zipfile

from kvk.rendering.new_source_export import ExportTable

WORKBOOK_VERSION = "reader_workbook_v1"
HEADER_ROW = 6
MAX_ROWS = 1_048_576
DIRECTORY = "Report Directory"


def column_layout(labels):
    """Bounded widths and wrapping shared by both physical workbook writers."""
    wide = {"Report", "Open report", "Source File", "Source"}
    prose = {
        "Purpose",
        "Notes",
        "Selection Notes",
        "Data Notes",
        "Issue",
        "Effect on Report",
        "Data Included",
        "Used In",
        "Applies To",
    }
    names = {
        "Player Name",
        "Camp",
        "Period",
        "Availability",
        "Data Status",
        "Import Status",
        "Source Sheet",
    }
    return tuple(
        (
            (
                380
                if label == "Source File"
                else (
                    340
                    if label in wide
                    else (
                        420
                        if label in prose
                        else 240 if label in names else 180 if "(UTC)" in label else 140
                    )
                )
            ),
            label in wide | prose | names,
        )
        for label in labels
    )


def cell_value(value):
    if value is None or value == "":
        return {}
    if isinstance(value, datetime):
        value = Decimal(str((value - datetime(1899, 12, 30, tzinfo=UTC)).total_seconds())) / 86400
    if type(value) in (int, Decimal):
        return {"numberValue": int(value) if value == int(value) else float(value)}
    if type(value) is not str:
        raise ValueError("Unsupported physical reader value.")
    if len(value.encode("utf-16-le")) // 2 > 32767 or re.search(
        r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value
    ):
        raise ValueError("Reader text exceeds the lossless Excel cell contract.")
    return {"stringValue": value}


def alias(title, used):
    clean = re.sub(r"[\x00-\x1f\[\]:*?/\\]", " ", title).strip(" '") or "Report"

    def truncate(text, units):
        while len(text.encode("utf-16-le")) // 2 > units:
            text = text[:-1]
        return text

    candidate, ordinal = truncate(clean, 31), 1
    while candidate.casefold() in used or candidate.casefold() == "history":
        ordinal += 1
        suffix = f" ({ordinal})"
        candidate = truncate(clean, 31 - len(suffix)) + suffix
    used.add(candidate.casefold())
    return candidate


def payload_hash(rows):
    def normalize(value):
        if isinstance(value, dict):
            return {k: normalize(v) for k, v in value.items() if v != {}}
        if isinstance(value, (list, tuple)):
            return [normalize(v) for v in value]
        if type(value) is float and value.is_integer():
            return int(value)
        return value

    return hashlib.sha256(
        json.dumps(
            normalize(rows),
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
            allow_nan=False,
        ).encode()
    ).hexdigest()


@dataclass(frozen=True)
class WorkbookTable:
    name: str
    logical_key: str
    values_json: str
    data_rows: int
    width: int
    frozen_columns: int

    def raw_values(self):
        return json.loads(self.values_json)

    @property
    def sha256(self):
        return payload_hash(self.raw_values())


@dataclass(frozen=True)
class WorkbookGeneration:
    projection: object
    tables: tuple[WorkbookTable, ...]

    @property
    def selections(self):
        return self.projection.audit.selections

    @property
    def key(self):
        return payload_hash([WORKBOOK_VERSION, self.projection.key, self.manifest()])

    def manifest(self):
        return {
            t.name: dict(
                version=WORKBOOK_VERSION,
                logical_key=t.logical_key,
                rows=t.data_rows,
                columns=t.width,
                sha256=t.sha256,
                frozen_columns=t.frozen_columns,
                column_widths=[
                    p
                    for p, _ in column_layout(
                        c.get("userEnteredValue", {}).get("stringValue", "")
                        for c in t.raw_values()[5]
                    )
                ],
                order=i,
            )
            for i, t in enumerate(self.tables)
        }


def build_workbook_generation(projection):
    used, tables = {DIRECTORY.casefold()}, []
    names = {s.key: alias(s.title, used) for s in projection.sheets}
    for sheet in projection.sheets:
        width = len(sheet.columns)
        rows = [[{} for _ in range(width)] for _ in range(5)]
        for i, text in enumerate(
            (sheet.title, sheet.period, sheet.purpose, sheet.availability, " | ".join(sheet.notes))
        ):
            rows[i][0] = {"userEnteredValue": cell_value(text)}
        rows.append(
            [
                {
                    "userEnteredValue": cell_value(c.label),
                    "userEnteredFormat": {
                        "textFormat": {"bold": True},
                        "backgroundColor": {"red": 0.85, "green": 0.9, "blue": 0.96},
                        "wrapStrategy": "WRAP",
                    },
                }
                for c in sheet.columns
            ]
        )
        layout = column_layout(c.label for c in sheet.columns)
        for row in sheet.rows:
            cells = []
            for column, value, (_, wrap) in zip(sheet.columns, row, layout, strict=True):
                fmt = {
                    "type": {"text": "TEXT", "number": "NUMBER", "datetime": "DATE_TIME"}[
                        column.kind
                    ],
                    "pattern": column.number_format,
                }
                # Excel displays a trailing decimal point for optional-only decimal
                # formats on integers. Keep fractional scores and explicit ratios intact.
                if (
                    column.kind == "number"
                    and column.number_format == "#,##0.######"
                    and value is not None
                    and value == int(value)
                ):
                    fmt["pattern"] = "#,##0"
                if column.number_format == "0.00%":
                    fmt["type"] = "PERCENT"
                cells.append(
                    {
                        "userEnteredValue": cell_value(value),
                        "userEnteredFormat": {"numberFormat": fmt},
                    }
                )
                if wrap:
                    cells[-1]["userEnteredFormat"]["wrapStrategy"] = "WRAP"
            rows.append(cells)
        for link in sheet.links:
            if "readerLink" in rows[link.row + HEADER_ROW][link.column]:
                raise ValueError(
                    "Reader cells require a single report link; use Report Links for multiple destinations."
                )
            rows[link.row + HEADER_ROW][link.column]["readerLink"] = names[link.target]
        identity = 2 if sheet.columns[0].label == "Governor ID" else 1
        tables.append(
            WorkbookTable(
                names[sheet.key],
                sheet.key,
                json.dumps(rows, ensure_ascii=False, separators=(",", ":")),
                len(sheet.rows),
                width,
                min(identity, width),
            )
        )
    return WorkbookGeneration(projection, tuple(tables))


def partition_workbook(manifest, *, max_cells=9_000_000, max_rows=MAX_ROWS):
    """Deterministic row slices, including banners and a complete first-file directory."""
    reserve = 10_000
    if type(max_cells) is not int or max_cells <= reserve or not HEADER_ROW < max_rows <= MAX_ROWS:
        raise ValueError("Invalid workbook capacity.")
    bins, current, cells, used = [], [], reserve, {DIRECTORY.casefold()}
    for name, spec in sorted(manifest.items(), key=lambda item: item[1]["order"]):
        width, total, offset = spec["columns"], spec["rows"], 0
        if (
            spec["version"] != WORKBOOK_VERSION
            or type(width) is not int
            or not 1 <= width <= 16384
            or type(total) is not int
            or total < 0
        ):
            raise ValueError("Invalid reader manifest dimensions/version.")
        while True:
            if len(current) >= (199 if not bins else 200):
                bins.append(current)
                current, cells = [], 0
            available = min(max_rows - HEADER_ROW, (max_cells - cells) // width - HEADER_ROW)
            if available < max(1, min(total - offset, 1)):
                if not current:
                    raise ValueError("Reader columns exceed workbook capacity.")
                bins.append(current)
                current, cells = [], 0
                continue
            count = min(total - offset, available)
            title = alias(name + (f" part {offset + 1}" if offset else ""), used)
            current.append(dict(table=name, name=title, offset=offset, rows=count, columns=width))
            cells += max(HEADER_ROW + 1, HEADER_ROW + count) * width
            offset += count
            if offset >= total:
                break
    if current:
        bins.append(current)
    if (sum(map(len, bins)) + HEADER_ROW) * 5 > reserve:
        raise ValueError("Reader directory exceeds reserved capacity.")
    return bins


def piece_values(table, piece):
    rows = table.raw_values()
    return (
        rows[:HEADER_ROW]
        + rows[HEADER_ROW + piece["offset"] : HEADER_ROW + piece["offset"] + piece["rows"]]
    )


def excel_bundle(generation, *, max_cells=9_000_000):
    """Return safe, fixed-name XLSX parts plus lossless audit and inert CSV in a ZIP.

    Uses the same typed physical tables and partitions as the Google transport.
    No filesystem writes or user-supplied archive paths occur here.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    layout = partition_workbook(generation.manifest(), max_cells=max_cells)
    tables = {t.name: t for t in generation.tables}
    targets = {
        p["table"]: (i, p["name"])
        for i, pieces in reversed(list(enumerate(layout)))
        for p in reversed(pieces)
    }
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for part, pieces in enumerate(layout):
            book = Workbook()
            book.remove(book.active)
            if part == 0:
                directory = book.create_sheet(DIRECTORY)
                directory.cell(1, 1, DIRECTORY)
                for col, label in enumerate(
                    ("Report", "Workbook", "First data row", "Rows", "Open report"), 1
                ):
                    directory.cell(6, col, label)
                for i, members in enumerate(layout):
                    for p in members:
                        directory.append([p["table"], i + 1, p["offset"] + 1, p["rows"], p["name"]])
                        for cell in directory[directory.max_row]:
                            if type(cell.value) is str:
                                cell.data_type = "s"
                        escaped = p["name"].replace("'", "''")
                        directory.cell(directory.max_row, 5).hyperlink = (
                            "#" if i == 0 else f"report-{i+1:02d}.xlsx#"
                        ) + f"'{escaped}'!A1"
                directory.freeze_panes = "B7"
                directory.auto_filter.ref = f"A6:E{directory.max_row}"
                for col in "ABCDE":
                    index = ord(col) - ord("A")
                    pixels, wrap = column_layout([c.value for c in directory[6]])[index]
                    directory.column_dimensions[col].width = (pixels - 5) / 7
                    for cell in directory[col][5:]:
                        cell.alignment = Alignment(wrap_text=wrap or cell.row == 6, vertical="top")
            for piece in pieces:
                table = tables[piece["table"]]
                ws = book.create_sheet(piece["name"])
                for r, values in enumerate(piece_values(table, piece), 1):
                    for c, spec in enumerate(values, 1):
                        value = spec.get("userEnteredValue", {})
                        cell = ws.cell(r, c, value.get("stringValue", value.get("numberValue")))
                        if "stringValue" in value:
                            cell.data_type = "s"
                        cell.number_format = (
                            spec.get("userEnteredFormat", {})
                            .get("numberFormat", {})
                            .get("pattern", "General")
                        )
                        if spec.get("userEnteredFormat", {}).get("wrapStrategy") == "WRAP":
                            cell.alignment = Alignment(wrap_text=True, vertical="top")
                        if "readerLink" in spec:
                            target_part, title = targets[spec["readerLink"]]
                            escaped = title.replace("'", "''")
                            cell.hyperlink = (
                                "#" if target_part == part else f"report-{target_part+1:02d}.xlsx#"
                            ) + f"'{escaped}'!A1"
                ws.freeze_panes = f"{get_column_letter(table.frozen_columns+1)}7"
                ws.auto_filter.ref = f"A6:{get_column_letter(table.width)}{max(7, piece['rows']+6)}"
                for cell in ws[6]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="17365D")
                    cell.alignment = Alignment(wrap_text=True, vertical="center")
                ws.row_dimensions[6].height = 48
                ws.cell(1, 1).font = Font(bold=True, size=16, color="17365D")
                for c in range(1, table.width + 1):
                    label = ws.cell(6, c).value
                    ws.column_dimensions[get_column_letter(c)].width = (
                        column_layout([label])[0][0] - 5
                    ) / 7
            content = io.BytesIO()
            book.save(content)
            archive.writestr(f"report-{part+1:02d}.xlsx", content.getvalue())
        audit = {
            t.name: dict(columns=t.columns, rows=t.rows) for t in generation.projection.audit.tables
        }
        archive.writestr("audit.json", json.dumps(audit, ensure_ascii=False))
        archive.writestr(
            "manifest.json",
            json.dumps(
                dict(
                    version=WORKBOOK_VERSION,
                    key=generation.key,
                    tables=generation.manifest(),
                    parts=layout,
                )
            ),
        )
        for i, sheet in enumerate(generation.projection.sheets, 1):
            archive.writestr(
                f"csv/report-{i:03d}.csv",
                ExportTable(
                    sheet.key, tuple(c.label for c in sheet.columns), sheet.exact_rows
                ).csv_bytes(),
            )
        archive.writestr(
            "reader-exact.json",
            json.dumps(
                {s.key: s.exact_rows for s in generation.projection.sheets}, ensure_ascii=False
            ),
        )
    return output.getvalue()
