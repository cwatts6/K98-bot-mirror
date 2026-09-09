"""Bounded, formula-preserving XLSX parsing into immutable private source facts."""

from contextlib import contextmanager
from dataclasses import replace
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation, localcontext
import hashlib
from io import BytesIO
import re
from xml.parsers import expat
from zipfile import ZIP_DEFLATED, BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.cell.read_only import EMPTY_CELL
from openpyxl.utils.cell import coordinate_to_tuple

from kvk.models.new_source_observation import (
    CampMapping,
    MetricValue,
    ParseDiagnostic,
    PreparedAggregateReport,
    PreparedPlayerObservation,
    SourceRow,
    TypedCell,
    ValidatedSourceMetadata,
)
from kvk.schemas.new_source_schema import (
    AGGREGATE_METRICS,
    AGGREGATE_SCHEMA_VERSION,
    BIGINT_MAX,
    CAMP_HEADERS,
    CAMP_SHEET,
    INT_MAX,
    KINGDOM_HEADERS,
    KINGDOM_SHEET,
    PLAYER_HEADERS,
    PLAYER_SCHEMA_VERSION,
    PLAYER_SHEET,
    PLAYER_TEXT_HEADERS,
    MetricState,
    ParseLimits,
    SourceKind,
    SourceValidationError,
    validate_headers,
)
from kvk.services.new_source_digest import canonical_decimal, semantic_digest_v1
from kvk.services.new_source_metadata import validate_source_metadata

_LITERAL = re.compile(r"[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[Ee][+-]?[0-9]+)?")
_REPORT_NUMBER = re.compile(
    r"(?P<number>(?:[0-9]+|[1-9][0-9]{0,2}(?:,[0-9]{3})+)(?:\.[0-9]+)?)" r"(?P<suffix>[KMBTkmbt]?)"
)


def _reject(code: str, message: str) -> None:
    raise SourceValidationError(code, message)


def _decimal(raw: str | None) -> Decimal | None:
    if raw is None or not _LITERAL.fullmatch(raw) or len(raw) > 128:
        return None
    try:
        number = Decimal(raw)
        # Bound pathological exponents before any fixed-point serialization/arithmetic.
        return number if number.is_finite() and abs(number.as_tuple().exponent) <= 1000 else None
    except InvalidOperation:
        return None


def _xml_preflight(data: bytes, limits: ParseLimits, total_cells: list[int]) -> dict[str, str]:
    """Expat rejects entities before expansion; count actual cells and preserve numeric lexemes."""
    parser = expat.ParserCreate(namespace_separator="}")
    numbers: dict[str, str] = {}
    seen_cells: set[str] = set()
    cell = None
    numeric = False
    in_value = False
    value_parts: list[str] = []
    text_length = 0
    depth = 0
    text_depth = 0
    text_count = 0
    previous_row = 0
    previous_column = 0

    def forbidden(*args):
        _reject("xml_declaration", "DTD and entity declarations are not accepted.")

    def start(name, attrs):
        nonlocal cell, numeric, in_value, text_length, depth, text_depth, text_count
        nonlocal previous_row, previous_column
        local = name.rsplit("}", 1)[-1]
        depth += 1
        if depth > 64 or any(len(v) > limits.max_cell_characters for v in attrs.values()):
            _reject("xml_limit", "XML nesting or attribute size exceeds parser limits.")
        if local == "Relationship" and attrs.get("TargetMode", "").lower() == "external":
            _reject("external_relationship", "External workbook relationships are disabled.")
        if local in ("Override", "Default") and any(
            "macroenabled" in v.lower() or "vbaproject" in v.lower() for v in attrs.values()
        ):
            _reject("macro_package", "Only macro-free XLSX packages are accepted.")
        if local in ("si", "is"):
            text_depth = depth
            text_count = 0
        if local == "row":
            row_token = attrs.get("r", "")
            if not re.fullmatch(r"[1-9][0-9]{0,6}", row_token) or int(row_token) <= previous_row:
                _reject("row_coordinate", "Worksheet rows must be ordered and uniquely numbered.")
            previous_row, previous_column = int(row_token), 0
        if local == "c":
            total_cells[0] += 1
            if total_cells[0] > limits.max_populated_cells:
                _reject("cell_limit", "Populated cell budget exceeded.")
            cell = attrs.get("r")
            if (
                cell is None
                or cell in seen_cells
                or not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,6}", cell)
            ):
                _reject("cell_coordinate", "Malformed or duplicate source cell coordinate.")
            seen_cells.add(cell)
            row, column = coordinate_to_tuple(cell)
            if row != previous_row or column <= previous_column:
                _reject("cell_order", "Worksheet cells must match their row in column order.")
            previous_column = column
            # Bound sparse iteration before openpyxl, independently of declared dimensions.
            if (
                column > limits.max_columns
                or row > max(limits.max_player_rows, limits.max_kingdom_rows, limits.max_camps) + 12
            ):
                _reject("worksheet_span", "Worksheet row/column span exceeds parser limits.")
            numeric = attrs.get("t", "n") == "n"
            text_length = 0
            value_parts.clear()
        if local == "v" and cell is not None:
            in_value = True

    def characters(text):
        nonlocal text_length, text_count
        if text_depth:
            text_count += len(text)
            if text_count > limits.max_cell_characters:
                _reject("cell_text_limit", "Source text exceeds the cell character limit.")
        if cell is not None:
            text_length += len(text)
            if text_length > limits.max_cell_characters:
                _reject("cell_text_limit", "Source text exceeds the cell character limit.")
            if numeric and in_value:
                value_parts.append(text)

    def end(name):
        nonlocal cell, in_value, depth, text_depth
        local = name.rsplit("}", 1)[-1]
        if local == "v":
            in_value = False
        if local == "c":
            if numeric and value_parts:
                numbers[cell] = "".join(value_parts)
            cell = None
        if depth == text_depth:
            text_depth = 0
        depth -= 1

    parser.StartElementHandler = start
    parser.EndElementHandler = end
    parser.CharacterDataHandler = characters
    parser.StartDoctypeDeclHandler = forbidden
    parser.EntityDeclHandler = forbidden
    parser.ExternalEntityRefHandler = forbidden
    parser.Parse(data, True)
    return numbers


@contextmanager
def _workbook(content: bytes, limits: ParseLimits):
    if not isinstance(content, bytes) or not content or len(content) > limits.max_compressed_bytes:
        _reject("compressed_size", "Supply nonempty XLSX bytes within the compressed size limit.")
    workbook = None
    try:
        numeric_cells = {}
        parsing_overrides = {}
        with BytesIO(content) as source, ZipFile(source) as archive:
            entries = archive.infolist()
            if len(entries) > limits.max_zip_members:
                _reject("zip_members", "ZIP member count exceeds the limit.")
            names = [entry.filename for entry in entries]
            if len(set(names)) != len(names):
                _reject("zip_duplicate", "Duplicate ZIP members are not accepted.")
            total_advertised = sum(e.file_size for e in entries)
            total_compressed = sum(e.compress_size for e in entries)
            if (
                total_advertised > limits.max_uncompressed_bytes
                or total_advertised > max(1, total_compressed) * limits.max_compression_ratio
            ):
                _reject("zip_total_size", "ZIP expansion exceeds the total size or ratio limit.")
            actual_total = 0
            total_cells = [0]
            for entry in entries:
                name = entry.filename
                if (
                    name.startswith("/")
                    or "\\" in name
                    or ":" in name
                    or any(p in ("", ".", "..") for p in name.split("/"))
                    or "\x00" in entry.orig_filename
                    or entry.flag_bits & 1
                    or entry.compress_type not in (0, 8)
                ):
                    _reject("zip_member", "Unsafe, encrypted or unsupported ZIP member.")
                if name.lower().endswith(".bin"):
                    _reject("macro_package", "Binary workbook parts are not accepted.")
                if (
                    entry.file_size > limits.max_entry_bytes
                    or entry.file_size > max(1, entry.compress_size) * limits.max_compression_ratio
                ):
                    _reject(
                        "zip_entry_size", "ZIP entry expansion exceeds the size or ratio limit."
                    )
                chunks = []
                actual_entry = 0
                with archive.open(entry) as stream:
                    while chunk := stream.read(64 * 1024):
                        actual_entry += len(chunk)
                        actual_total += len(chunk)
                        if (
                            actual_entry > limits.max_entry_bytes
                            or actual_total > limits.max_uncompressed_bytes
                            or actual_entry
                            > max(1, entry.compress_size) * limits.max_compression_ratio
                            or actual_total
                            > max(1, total_compressed) * limits.max_compression_ratio
                        ):
                            _reject(
                                "zip_actual_size", "Actual ZIP expansion exceeds parser limits."
                            )
                        # Workbook relationships, not suffixes, decide what the
                        # loader consumes. Preflight every part in this XML-only
                        # source schema, including renamed worksheet parts.
                        chunks.append(chunk)
                if actual_entry != entry.file_size:
                    _reject("zip_size_mismatch", "ZIP entry size is inconsistent.")
                if chunks:
                    xml = b"".join(chunks)
                    values = _xml_preflight(xml, limits, total_cells)
                    if values:
                        numeric_cells[name] = values
                        invalid = {key for key, value in values.items() if _decimal(value) is None}
                        if invalid:
                            # openpyxl casts numeric XML before handing cells to us. A
                            # bounded private parsing view prevents invalid optional
                            # numbers from aborting the whole player file. Original
                            # tokens below remain authoritative, never this placeholder.
                            encoding = (
                                "utf-16" if xml.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8"
                            )
                            text = xml.decode(encoding)

                            def replace_invalid(match):
                                coordinate = re.search(r"\br=['\"]([A-Z]+[0-9]+)['\"]", match[1])
                                if coordinate is None or coordinate[1] not in invalid:
                                    return match[0]
                                body = re.sub(
                                    r"(<(?:[A-Za-z_][\w.-]*:)?v>)[^<]*(</(?:[A-Za-z_][\w.-]*:)?v>)",
                                    r"\g<1>0\g<2>",
                                    match[2],
                                )
                                return match[1] + body + match[3]

                            text = re.sub(
                                r"(<(?:[A-Za-z_][\w.-]*:)?c\b[^>]*>)(.*?)(</(?:[A-Za-z_][\w.-]*:)?c>)",
                                replace_invalid,
                                text,
                                flags=re.DOTALL,
                            )
                            parsing_overrides[name] = text.encode(encoding)
        parsing_content = content
        if parsing_overrides:
            with BytesIO(content) as source, ZipFile(source) as original, BytesIO() as output:
                with ZipFile(output, "w", compression=ZIP_DEFLATED) as view:
                    for entry in original.infolist():
                        data = parsing_overrides.get(entry.filename)
                        view.writestr(
                            entry.filename, original.read(entry) if data is None else data
                        )
                parsing_content = output.getvalue()
        with BytesIO(parsing_content) as source:
            workbook = load_workbook(source, read_only=True, data_only=False, keep_links=False)
            yield workbook, numeric_cells
    except SourceValidationError:
        raise
    except (
        BadZipFile,
        expat.ExpatError,
        KeyError,
        ValueError,
        TypeError,
        OSError,
        OverflowError,
        RuntimeError,
        NotImplementedError,
    ) as exc:
        raise SourceValidationError("invalid_xlsx", "Corrupt or unsupported XLSX package.") from exc
    finally:
        if workbook is not None:
            workbook.close()


def _cell(header, cell, numeric_cells, limits) -> TypedCell:
    value = cell.value
    tag = cell.data_type
    numeric = None
    if value is None or value == "":
        tag, raw = "null", None
    elif tag == "f":
        if isinstance(value, str):
            raw = value
        else:
            # Array/data-table descriptors are not literal formula strings. Do
            # not hash their process-specific object repr as source evidence.
            _reject("unsupported_formula", "Nonliteral formula descriptors require schema review.")
    elif isinstance(value, bool):
        tag, raw = "b", "true" if value else "false"
    elif isinstance(value, (datetime, date, time)):
        tag, raw = "d", value.isoformat()
    elif tag == "n":
        raw = numeric_cells.get(cell.coordinate, str(value))
        numeric = _decimal(raw)
    else:
        tag = "e" if tag == "e" else "s"
        raw = str(value)
    if raw is not None and len(raw) > limits.max_cell_characters:
        _reject("cell_text_limit", "Source text exceeds the cell character limit.")
    invalid = MetricValue(None, MetricState.INVALID_SOURCE_VALUE, raw)
    return TypedCell(header, tag, raw, numeric, str(cell.number_format), invalid)


def _integer(cell: TypedCell, *, identity=False, maximum=BIGINT_MAX) -> Decimal | None:
    value = (
        cell.numeric_value
        if cell.cell_type == "n"
        else (
            _decimal(cell.raw_value)
            if cell.cell_type == "s"
            and cell.raw_value is not None
            and re.fullmatch(r"[+]?[0-9]+", cell.raw_value)
            else None
        )
    )
    if (
        value is None
        or value != value.to_integral_value()
        or not (1 if identity else 0) <= value <= maximum
    ):
        return None
    # Excel's numeric identity precision is only 15 significant decimal digits.
    if identity and cell.cell_type == "n" and value > 999_999_999_999_999:
        return None
    return value


def _player_cell(cell: TypedCell) -> TypedCell:
    if cell.header in PLAYER_TEXT_HEADERS:
        value = cell.raw_value if cell.cell_type == "s" and len(cell.raw_value) <= 256 else None
    else:
        value = _integer(
            cell,
            identity=cell.header in ("Governor ID", "Kingdom"),
            maximum=INT_MAX if cell.header == "Kingdom" else BIGINT_MAX,
        )
    return replace(
        cell,
        metric=MetricValue(
            value,
            MetricState.AVAILABLE if value is not None else MetricState.INVALID_SOURCE_VALUE,
            cell.raw_value,
        ),
    )


def _aggregate_cell(cell: TypedCell) -> TypedCell:
    if cell.raw_value is None or len(cell.raw_value) > 128 or cell.cell_type not in ("s", "n"):
        _reject("aggregate_metric", "Every required aggregate metric must be valid.")
    raw = (
        canonical_decimal(cell.numeric_value) if cell.numeric_value is not None else cell.raw_value
    )
    match = _REPORT_NUMBER.fullmatch(raw)
    if not match:
        _reject("aggregate_metric", "Aggregate metrics require strict nonnegative decimal tokens.")
    token = match["number"].replace(",", "")
    suffix = match["suffix"].upper()
    power = {"": 0, "K": 3, "M": 6, "B": 9, "T": 12}[suffix]
    fractional_digits = len(token.split(".")[1]) if "." in token else 0
    with localcontext() as context:
        context.prec = 160
        value = Decimal(token) * Decimal(10) ** power
        unit = Decimal(10) ** (power - fractional_digits)
        if value >= Decimal(10) ** 32 or value != value.quantize(Decimal("0.000001")):
            _reject("aggregate_decimal_range", "Aggregate metric cannot fit decimal(38,6) exactly.")
        if unit < Decimal("0.000001"):
            _reject("aggregate_precision", "Displayed precision cannot fit decimal(38,6).")
    metric = MetricValue(
        value,
        MetricState.AVAILABLE,
        cell.raw_value,
        unit,
        "reported_abbreviated" if suffix else "reported_numeric",
    )
    return replace(cell, metric=metric)


def _table(sheet, expected, numeric_cells, limits, max_rows, *, tail=False):
    sheet.reset_dimensions()
    iterator = sheet.iter_rows()
    header_row = next(iterator, ())
    headers = tuple(c.value for c in header_row)
    while headers and headers[-1] is None:
        headers = headers[:-1]
    validate_headers(headers, expected)
    indices = {header: i for i, header in enumerate(headers)}
    rows = []
    diagnostics = []
    separated = False
    tail_coordinates = []
    for number, row in enumerate(iterator, 2):
        populated = [i for i, c in enumerate(row) if c.value is not None and c.value != ""]
        if not populated:
            separated = True
            continue
        if any(i >= len(headers) for i in populated):
            _reject("unknown_column", "Populated unknown columns require schema review.")
        if separated:
            # Only the documented A1 shape is approved: at most nine numeric E-column notes.
            if (
                not tail
                or populated != [4]
                or len(tail_coordinates) >= 9
                or row[4].data_type != "n"
                or _decimal(str(row[4].value)) is None
            ):
                _reject("novel_tail", "Keyed rows or novel populated tail content require review.")
            tail_coordinates.append(f"E{number}")
            continue
        if len(rows) >= max_rows:
            _reject("row_limit", "Keyed row count exceeds the configured parser limit.")
        cells = tuple(
            _cell(
                h, row[indices[h]] if indices[h] < len(row) else EMPTY_CELL, numeric_cells, limits
            )
            for h in expected
        )
        rows.append(cells)
    if not rows:
        _reject("empty_table", "A recognized keyed table must contain rows.")
    if tail_coordinates:
        diagnostics.append(
            ParseDiagnostic(
                "ignored_tail_annotations",
                sheet.title,
                len(tail_coordinates),
                f"{tail_coordinates[0]}:{tail_coordinates[-1]}",
            )
        )
    return rows, diagnostics


def _metadata(metadata, kind):
    if not isinstance(metadata, ValidatedSourceMetadata) or metadata.candidate.kind != kind:
        _reject("metadata_kind", "Workbook kind must match validated metadata.")
    # Do not let direct construction of a frozen dataclass bypass validation.
    checked = validate_source_metadata(metadata.candidate, metadata.confirmation, metadata.scope)
    if checked != metadata:
        _reject(
            "metadata_confirmation_mismatch", "Metadata must retain its validated confirmation."
        )


def parse_player_workbook(
    content: bytes,
    metadata: ValidatedSourceMetadata,
    limits: ParseLimits = ParseLimits(),
) -> PreparedPlayerObservation:
    _metadata(metadata, SourceKind.PLAYERS)
    with _workbook(content, limits) as (workbook, numeric_cells):
        if workbook.sheetnames != [PLAYER_SHEET]:
            _reject("schema_sheets", "Player XLSX requires only the Scan tab.")
        sheet = workbook[PLAYER_SHEET]
        raw_rows, diagnostics = _table(
            sheet,
            PLAYER_HEADERS,
            numeric_cells.get(sheet._worksheet_path, {}),
            limits,
            limits.max_player_rows,
        )
        rows = []
        seen = set()
        unavailable = 0
        for raw_row in raw_rows:
            cells = tuple(_player_cell(c) for c in raw_row)
            key = cells[0].metric.value
            kingdom = cells[-1].metric.value
            if key is None or kingdom is None or int(kingdom) not in metadata.scope.kingdoms:
                _reject("invalid_identity", "Player identity or configured kingdom is invalid.")
            if key in seen:
                _reject("duplicate_identity", "Duplicate governor identities reject the workbook.")
            seen.add(key)
            unavailable += sum(c.metric.state != MetricState.AVAILABLE for c in cells)
            rows.append(SourceRow(PLAYER_SHEET, int(key), cells, kingdom=int(kingdom)))
        if unavailable:
            diagnostics.append(
                ParseDiagnostic("unavailable_player_fields", PLAYER_SHEET, unavailable)
            )
        immutable_rows = tuple(sorted(rows, key=lambda r: r.key))
        return PreparedPlayerObservation(
            metadata,
            immutable_rows,
            hashlib.sha256(content).hexdigest(),
            PLAYER_SCHEMA_VERSION,
            semantic_digest_v1(immutable_rows, PLAYER_SCHEMA_VERSION),
            tuple(diagnostics),
        )


def _camp_key(value: str) -> str:
    return " ".join(value.split()).casefold()


def _mapping(mapping, scope, limits):
    if (
        not isinstance(mapping, CampMapping)
        or not isinstance(mapping.entries, tuple)
        or not mapping.entries
        or len(mapping.entries) > limits.max_kingdom_rows
    ):
        _reject("invalid_mapping", "An immutable complete camp mapping is required.")
    kingdoms, camps, names = {}, {}, {}
    for entry in mapping.entries:
        if not isinstance(entry, tuple) or len(entry) != 3:
            _reject("invalid_mapping", "Malformed camp mapping entry.")
        kingdom, camp, name = entry
        if (
            type(kingdom) is not int
            or kingdom not in scope.kingdoms
            or kingdom in kingdoms
            or type(camp) is not int
            or not 1 <= camp <= 8
            or not isinstance(name, str)
            or not 0 < len(name) <= 256
            or not _camp_key(name)
        ):
            _reject("invalid_mapping", "Camp mapping identities or labels are invalid.")
        normalized = _camp_key(name)
        if (camp in camps and camps[camp] != normalized) or (
            normalized in names and names[normalized] != camp
        ):
            _reject("ambiguous_camp", "Camp names must resolve uniquely within the mapping.")
        kingdoms[kingdom], camps[camp], names[normalized] = camp, normalized, camp
    if set(kingdoms) != set(scope.kingdoms) or len(camps) > limits.max_camps:
        _reject("mapping_scope", "Camp mapping must cover the configured scope exactly.")
    return kingdoms, camps, names


def parse_aggregate_workbook(
    content: bytes,
    metadata: ValidatedSourceMetadata,
    mapping: CampMapping,
    limits: ParseLimits = ParseLimits(),
) -> PreparedAggregateReport:
    _metadata(metadata, SourceKind.AGGREGATE)
    kingdoms, camps, names = _mapping(mapping, metadata.scope, limits)
    results, diagnostics = {}, []
    with _workbook(content, limits) as (workbook, numeric_cells):
        if set(workbook.sheetnames) != {KINGDOM_SHEET, CAMP_SHEET}:
            _reject("schema_sheets", "Aggregate XLSX requires exactly both aggregate tabs.")
        for title, headers, maximum in (
            (KINGDOM_SHEET, KINGDOM_HEADERS, limits.max_kingdom_rows),
            (CAMP_SHEET, CAMP_HEADERS, limits.max_camps),
        ):
            sheet = workbook[title]
            raw_rows, notes = _table(
                sheet,
                headers,
                numeric_cells.get(sheet._worksheet_path, {}),
                limits,
                maximum,
                tail=title == KINGDOM_SHEET,
            )
            diagnostics.extend(notes)
            rows, seen = [], set()
            for raw_row in raw_rows:
                label = raw_row[0]
                if label.cell_type != "s" or len(label.raw_value) > 256:
                    _reject("camp_identity", "Camp identity must be bounded literal text.")
                camp = names.get(_camp_key(label.raw_value))
                if camp is None:
                    _reject("unknown_camp", "Camp name is not in the configured mapping.")
                kingdom = None
                if title == KINGDOM_SHEET:
                    number = _integer(raw_row[1], identity=True, maximum=INT_MAX)
                    kingdom = int(number) if number is not None else None
                    if kingdom not in kingdoms or kingdoms[kingdom] != camp:
                        _reject("kingdom_camp_conflict", "Kingdom and camp must match the mapping.")
                key = kingdom if kingdom is not None else camp
                if key in seen:
                    _reject(
                        "duplicate_identity", "Duplicate aggregate identities reject both tabs."
                    )
                seen.add(key)
                cells = tuple(
                    (
                        _aggregate_cell(c)
                        if c.header in AGGREGATE_METRICS
                        else replace(
                            c,
                            metric=MetricValue(
                                Decimal(kingdom) if c.header == "KD" else c.raw_value,
                                MetricState.AVAILABLE,
                                c.raw_value,
                            ),
                        )
                    )
                    for c in raw_row
                )
                rows.append(SourceRow(title, key, cells, kingdom=kingdom, camp_id=camp))
            if seen != set(kingdoms if title == KINGDOM_SHEET else camps):
                _reject("aggregate_scope", "Both tabs must contain every configured identity once.")
            results[title] = tuple(sorted(rows, key=lambda r: r.key))
        all_rows = results[KINGDOM_SHEET] + results[CAMP_SHEET]
        return PreparedAggregateReport(
            metadata,
            results[KINGDOM_SHEET],
            results[CAMP_SHEET],
            mapping,
            hashlib.sha256(content).hexdigest(),
            AGGREGATE_SCHEMA_VERSION,
            semantic_digest_v1(all_rows, AGGREGATE_SCHEMA_VERSION),
            tuple(diagnostics),
        )
